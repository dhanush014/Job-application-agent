"""Write the same resume as an editable .docx.

Built from the ResumeDoc the PDF was rendered from, so the two always carry
identical content. Deliberately plain: no italics, no tables, no text boxes,
no headers or footers, because applicant tracking systems parse those poorly.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt, RGBColor

from ..models import DocEntry, ResumeDoc

RULE = "—" * 40  # em-dashes under a section heading


def _base(doc: Document, family: str, pt: float) -> None:
    style = doc.styles["Normal"]
    style.font.name = family
    style.font.size = Pt(pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(36)
        section.left_margin = section.right_margin = Pt(36)


def _para(doc: Document, *, before: float = 0, after: float = 0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def _run(p, text: str, *, bold: bool = False, size: float | None = None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = False
    if size:
        r.font.size = Pt(size)
    return r


def _heading(doc: Document, title: str, pt: float) -> None:
    p = _para(doc, before=pt * 0.9, after=pt * 0.15)
    _run(p, title.upper(), bold=True, size=pt + 0.5)
    rule = _para(doc, after=pt * 0.1)
    _run(rule, RULE, size=pt * 0.5)


def _entry_header(doc: Document, left: str, right: str, pt: float, *, bold_left: bool, before: float) -> None:
    """Left text with the date right-aligned on the same line, via a tab stop."""
    p = _para(doc, before=before)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(540), WD_TAB_ALIGNMENT.RIGHT)
    _run(p, left, bold=bold_left)
    if right:
        _run(p, "\t" + right)


def _entry(doc: Document, e: DocEntry, pt: float) -> None:
    _entry_header(doc, e.heading, e.dates, pt, bold_left=True, before=pt * 0.55)
    if e.subheading or e.location:
        _entry_header(doc, e.subheading, e.location, pt, bold_left=False, before=0)
    for i, bullet in enumerate(e.bullets):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(pt * 0.18)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        segments = e.segments[i] if len(e.segments) > i else []
        if segments:
            for seg in segments:
                _run(p, seg.t, bold=seg.b)
        else:
            _run(p, bullet)


def write_docx(doc_model: ResumeDoc, path: Path) -> Path:
    pt = doc_model.layout.point_size
    family = doc_model.layout.family
    d = Document()
    _base(d, family, pt)

    name = _para(d, after=pt * 0.1)
    name.alignment = 1  # centre
    _run(name, doc_model.contact.name, bold=True, size=pt * 1.7)

    contact = _para(d, after=0)
    contact.alignment = 1
    bits = [doc_model.contact.location, doc_model.contact.phone, doc_model.contact.email]
    bits += [l.label for l in doc_model.contact.links]
    _run(contact, "  |  ".join(b for b in bits if b), size=pt - 0.5)

    _heading(d, "Education", pt)
    for e in doc_model.education:
        dates = f"{e.start} – {e.end}" if e.start else e.end
        _entry_header(d, e.school, dates, pt, bold_left=True, before=pt * 0.55)
        if e.degree or e.location:
            _entry_header(d, e.degree, e.location, pt, bold_left=False, before=0)
        for detail in e.details:
            p = d.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(pt * 0.18)
            p.paragraph_format.space_after = Pt(0)
            _run(p, detail)

    for title, entries in (
        ("Experience", doc_model.roles),
        ("Academic Experience", doc_model.academic),
        ("Projects", doc_model.projects),
    ):
        if not entries:
            continue
        _heading(d, title, pt)
        for e in entries:
            _entry(d, e, pt)

    _heading(d, "Skills", pt)
    for line in doc_model.skills:
        p = _para(d, before=pt * 0.2)
        _run(p, line.category + ": ", bold=True)
        if line.segments:
            for seg in line.segments:
                _run(p, seg.t, bold=seg.b)
        else:
            _run(p, line.items)

    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path
