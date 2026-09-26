"""Excel export of every application (resume, JD, answers, dates)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..models import Application

COLUMNS = [
    ("Applied at", 20), ("Status", 13), ("Company", 20), ("Title", 34), ("Location", 20),
    ("Score", 7), ("Fit summary", 50), ("Job URL", 40), ("Resume PDF", 40),
    ("Resume bullets", 70), ("Answers", 80), ("Cover letter", 50), ("Review reasons", 40),
    ("Job description", 80), ("Attempt", 8), ("Created at", 20), ("Error", 30),
]
EXCEL_CELL_LIMIT = 32_000


def _cell(s) -> str:
    s = "" if s is None else str(s)
    return s[:EXCEL_CELL_LIMIT]


def _bullets(app: Application) -> str:
    doc = app.resume_json or {}
    out = []
    for e in doc.get("roles", []) + doc.get("academic", []) + doc.get("projects", []):
        out.append(e.get("heading", ""))
        out += [f"  • {b}" for b in e.get("bullets", [])]
    return "\n".join(out)


def _answers(app: Application) -> str:
    lines = []
    for a in app.answers:
        v = ", ".join(a.value) if isinstance(a.value, list) else a.value
        lines.append(f"Q: {a.label}\nA: {v if v is not None else '—'}  [{a.source}]")
    return "\n\n".join(lines)


def export_xlsx(apps: list[Application], path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append([c for c, _ in COLUMNS])
    for i, (_, w) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True)
    ws.freeze_panes = "A2"
    for a in sorted(apps, key=lambda a: (a.applied_at or a.created_at), reverse=True):
        ws.append([
            _cell(a.applied_at.strftime("%Y-%m-%d %H:%M") if a.applied_at else ""),
            a.status.value, _cell(a.company_name), _cell(a.title), _cell(a.location), a.score,
            _cell(a.fit_summary), _cell(a.url), _cell(a.resume_pdf), _cell(_bullets(a)),
            _cell(_answers(a)), _cell(a.cover_letter), _cell("\n".join(a.review_reasons)),
            _cell(a.jd), a.attempt, _cell(a.created_at.strftime("%Y-%m-%d %H:%M")), _cell(a.error),
        ])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
