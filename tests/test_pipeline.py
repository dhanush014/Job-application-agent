from datetime import datetime, timedelta, timezone
from pathlib import Path

from conftest import FakeLLM, mock_http

from jobagent.apply.browser import SubmitResult
from jobagent.models import Status
from jobagent.pipeline import Pipeline
from jobagent.store.export import export_xlsx


class FakeSubmitter:
    def __init__(self, outcome="submitted"):
        self.outcome = outcome
        self.calls = []

    def submit(self, app, files, *, headless, dry_run):
        self.calls.append((app.id, headless, dry_run))
        assert files.resume.exists()
        if dry_run:
            return SubmitResult("dry_run", "filled")
        return SubmitResult(self.outcome, "ok" if self.outcome == "submitted" else "CAPTCHA challenge on the page")


def make(cfg, store, **kw):
    sub = kw.pop("submitter", FakeSubmitter())
    return Pipeline(cfg, store, kw.pop("llm", FakeLLM()), http=mock_http(), submitter=sub), sub


def test_full_run_dry_run_mode(cfg, store):
    p, sub = make(cfg, store)
    stats = p.run()
    assert stats.discovered_new == 3
    assert stats.filtered == 1  # "Engineering Manager" excluded by title
    apps = {a.title: a for a in store.list()}
    gh = apps["Senior Software Engineer, Payments Platform"]
    # clearance question answered "No" from the bank, everything else answered -> ready -> dry run
    assert gh.status == Status.DRY_RUN, gh.review_reasons
    assert gh.resume_pdf and gh.resume_json["roles"]
    assert any(a.field == "question_1004" and a.value for a in gh.answers)
    assert sub.calls and all(dry for _, _, dry in sub.calls)
    assert cfg.path(cfg.storage.excel_path).exists()


def test_live_submit_respects_daily_limit_and_records_date(cfg, store):
    cfg.apply.dry_run = False
    cfg.apply.daily_limit = 1
    p, sub = make(cfg, store)
    stats = p.run()
    assert stats.submitted == 1 and len(sub.calls) == 1
    applied = store.list([Status.APPLIED])
    assert len(applied) == 1 and applied[0].applied_at is not None
    assert p.run().submitted == 0  # limit reached for today


def test_browser_failure_lands_in_review_not_lost(cfg, store):
    cfg.apply.dry_run = False
    p, _ = make(cfg, store, submitter=FakeSubmitter("needs_human"))
    p.run()
    review = store.list([Status.NEEDS_REVIEW])
    assert review and any("CAPTCHA" in r for a in review for r in a.review_reasons)


def test_low_fit_is_not_prepared(cfg, store):
    p, sub = make(cfg, store, llm=FakeLLM(score=30))
    stats = p.run()
    assert stats.low_fit == 2 and not sub.calls


def test_reapply_after_cooldown(cfg, store):
    cfg.apply.dry_run = False
    p, _ = make(cfg, store)
    p.run()
    a = store.list([Status.APPLIED])[0]
    assert all(j.key != a.job_key for j, _ in store.jobs_to_process(True, 30))
    a.applied_at = datetime.now(timezone.utc) - timedelta(days=31)
    store.save(a)
    due = [(j, n) for j, n in store.jobs_to_process(True, 30) if j.key == a.job_key]
    assert due and due[0][1] == 2
    assert not [j for j, _ in store.jobs_to_process(False, 30) if j.key == a.job_key]


def test_manual_reapply_creates_new_attempt(cfg, store):
    cfg.apply.dry_run = False
    p, _ = make(cfg, store)
    p.run()
    a = store.list([Status.APPLIED])[0]
    b = p.reapply(a.id)
    assert b.id != a.id and b.attempt == 2 and b.status in (Status.READY, Status.NEEDS_REVIEW)


def test_excel_export_contains_answers(cfg, store, tmp_path):
    p, _ = make(cfg, store)
    p.run()
    from openpyxl import load_workbook

    path = export_xlsx(store.list(), tmp_path / "x.xlsx")
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert {"Applied at", "Resume PDF", "Answers", "Job description"} <= set(header)
    answers_col = header.index("Answers")
    assert any(r[answers_col].value and "authorized" in r[answers_col].value for r in ws.iter_rows(min_row=2))


def test_stuck_applying_is_recovered(cfg, store):
    p, _ = make(cfg, store)
    p.run(submit=False)
    a = store.list([Status.READY])[0]
    a.status = Status.APPLYING
    store.save(a)
    assert p.recover_stuck(older_than_minutes=30) == 0  # fresh, still in flight
    assert p.recover_stuck(older_than_minutes=-1) == 1
    assert store.get(a.id).status == Status.NEEDS_REVIEW


def test_cover_letter_written_rendered_and_attached(cfg, store):
    from pathlib import Path

    from pypdf import PdfReader

    p, _ = make(cfg, store)
    p.run(submit=False)
    gh = next(a for a in store.list() if a.job_id == "5001")
    assert gh.cover_letter and gh.cover_letter.startswith("Dear Hiring Team")
    assert Path(gh.cover_letter_pdf).exists()
    pdf = PdfReader(gh.cover_letter_pdf)
    text = pdf.pages[0].extract_text()
    assert len(pdf.pages) == 1 and "Alex Rivera" in text and "Acme" in text
    assert next(a for a in gh.answers if a.field == "cover_letter").value == "cover_letter"


def test_try_two_specific_jobs_dry_run_by_default(cfg, store):
    p, sub = make(cfg, store)
    results = p.try_urls([
        "https://job-boards.greenhouse.io/acme/jobs/5001",
        "https://jobs.ashbyhq.com/ramp/0b1c2d3e-aaaa-bbbb-cccc-000000000001",
    ])
    assert len(results) == 2 and len(sub.calls) == 2
    assert all(dry for _, headless, dry in sub.calls)  # nothing submitted
    assert all(not headless for _, headless, _ in sub.calls)  # visible browser for a trial
    for app, _ in results:
        assert app.resume_pdf and app.status in (Status.DRY_RUN, Status.NEEDS_REVIEW)


def test_try_live_submits_and_reports_filters(cfg, store):
    p, sub = make(cfg, store)
    (app, filtered), = p.try_urls(["https://job-boards.greenhouse.io/acme/jobs/5002"], live=True)
    assert filtered  # "Engineering Manager" is normally skipped by the title filters
    assert sub.calls[-1][2] is False and app.status == Status.APPLIED


def test_run_limit_caps_a_trial(cfg, store):
    cfg.apply.dry_run = False
    p, sub = make(cfg, store)
    stats = p.run(limit=1)
    assert stats.ready + stats.needs_review == 1  # only one job prepared
    assert len(sub.calls) <= 1


def test_shortlist_builds_a_folder_with_links_and_resumes(cfg, store, tmp_path):
    p, sub = make(cfg, store)
    out = tmp_path / "shortlist"
    index, apps = p.shortlist(limit=5, out_dir=out)
    assert not sub.calls, "shortlist must never open a browser"
    assert apps and index.exists()
    page = index.read_text()
    for a in apps:
        assert a.apply_url in page              # the link to apply
        assert a.title in page                  # what the job is
    pdfs = sorted(q.name for q in out.glob("*.pdf"))
    assert len(pdfs) >= len(apps)               # a tailored resume per job
    assert all(p.startswith("alex_rivera_resume_") for p in pdfs)  # easy to save
    assert sorted(q.name for q in out.glob("*.docx"))              # editable copy too
    assert all(Path(a.resume_pdf).exists() and Path(a.resume_docx).exists() for a in apps)
    # best-scoring job appears first in the page, whatever the filenames are
    assert page.index(apps[0].title) <= min(page.index(a.title) for a in apps)


def test_shortlist_from_specific_links(cfg, store, tmp_path):
    p, sub = make(cfg, store)
    index, apps = p.shortlist(
        urls=["https://job-boards.greenhouse.io/acme/jobs/5001"], limit=5, out_dir=tmp_path / "s"
    )
    assert len(apps) == 1 and not sub.calls
    assert "5001" in index.read_text()


def test_shortlist_skips_low_fit_jobs(cfg, store, tmp_path):
    p, _ = make(cfg, store, llm=FakeLLM(score=30))
    index, apps = p.shortlist(limit=5, out_dir=tmp_path / "s")
    assert apps == []
    assert "No jobs matched" in index.read_text()


def test_resume_font_setting_reaches_the_pdf(cfg, store):
    cfg.resume.font = "New Computer Modern"
    cfg.resume.min_pt, cfg.resume.max_pt = 9.5, 11.0
    p, _ = make(cfg, store)
    p.run(submit=False)
    app = next(a for a in store.list() if a.resume_json)
    layout = app.resume_json["layout"]
    assert layout["family"] == "New Computer Modern"
    assert 9.5 <= layout["min_pt"] and layout["max_pt"] <= 11.0
